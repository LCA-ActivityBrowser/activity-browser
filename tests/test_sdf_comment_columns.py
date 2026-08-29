"""SDF comments and required flow-scenario headers."""
from openpyxl import Workbook

from activity_browser.bwutils.superstructure.excel import (
    import_from_excel,
    valid_cols,
)
from activity_browser.bwutils.superstructure.file_imports import ABCSVImporter
from activity_browser.bwutils.superstructure.utils import (
    SUPERSTRUCTURE,
    is_flow_sdf_headers,
    is_partial_flow_sdf_headers,
    missing_superstructure_columns,
)


def _sample_row(amount: float = 1.0) -> list:
    base = {
        "from activity name": "A",
        "from reference product": "p",
        "from location": "GLO",
        "from categories": "",
        "from database": "db1",
        "from key": "('db1', 'a')",
        "to activity name": "B",
        "to reference product": "q",
        "to location": "GLO",
        "to categories": "",
        "to database": "db2",
        "to key": "('db2', 'b')",
        "flow type": "technosphere",
    }
    return [base[c] for c in SUPERSTRUCTURE] + ["note text", amount]


def _write_xlsx(path, rows) -> None:
    wb = Workbook()
    info = wb.active
    info.title = "info"
    info["A1"] = "info"
    sheet = wb.create_sheet("scenarios")
    for r_i, row in enumerate(rows, start=1):
        for c_i, val in enumerate(row, start=1):
            sheet.cell(r_i, c_i, val)
    wb.save(path)


def test_valid_cols_drops_underscore_prefix():
    assert valid_cols("_notes") is False
    assert valid_cols("2025") is True
    assert valid_cols("from database") is True


def test_missing_superstructure_columns_detects_typo():
    cols = [c if c != "to key" else "to keyhhh" for c in SUPERSTRUCTURE] + ["2025"]
    assert missing_superstructure_columns(cols) == ["to key"]
    assert is_flow_sdf_headers(cols) is False
    assert is_partial_flow_sdf_headers(cols) is True
    assert is_flow_sdf_headers(list(SUPERSTRUCTURE) + ["2025"]) is True


def test_csv_hash_rows_and_underscore_columns(tmp_path):
    cols = list(SUPERSTRUCTURE) + ["_notes", "2025"]
    header = ";".join(cols)
    data = (
        "A;p;GLO;;db1;('db1', 'a');B;q;GLO;;db2;('db2', 'b');technosphere;x;1.0"
    )
    text = "\n".join(
        [
            "# leading file comment",
            header,
            data,
            "# skipped data row",
        ]
    )
    path = tmp_path / "sdf.csv"
    path.write_text(text, encoding="utf-8")

    df = ABCSVImporter.read_file(path, separator=";")

    assert list(df.columns) == cols[:-2] + ["2025"]  # _notes dropped
    assert "_notes" not in df.columns
    assert len(df) == 1
    assert df.loc[0, "from database"] == "db1"
    assert float(df.loc[0, "2025"]) == 1.0


def test_csv_misspelled_header_is_partial_sdf(tmp_path):
    cols = [c if c != "to key" else "to keyhhh" for c in SUPERSTRUCTURE] + ["2025"]
    header = ";".join(cols)
    data = "A;p;GLO;;db1;('db1', 'a');B;q;GLO;;db2;('db2', 'b');technosphere;1.0"
    path = tmp_path / "bad_header.csv"
    path.write_text(header + "\n" + data + "\n", encoding="utf-8")

    df = ABCSVImporter.read_file(path, separator=";")
    assert is_partial_flow_sdf_headers(df.columns)
    assert missing_superstructure_columns(df.columns) == ["to key"]


def test_excel_leading_and_mid_hash_rows_with_underscore_column(tmp_path):
    """Leading/mid '#' rows and '_' columns are excluded via skiprows/usecols."""
    cols = list(SUPERSTRUCTURE) + ["_notes", "2025"]
    rows = [
        ["# leading file comment"] + [None] * (len(cols) - 1),
        cols,
        _sample_row(1.0),
        ["# skipped data row"] + [None] * (len(cols) - 1),
        _sample_row(2.0),
    ]
    path = tmp_path / "sdf.xlsx"
    _write_xlsx(path, rows)

    df = import_from_excel(path, import_sheet=1)

    assert not df.empty
    assert "from activity name" in df.columns
    assert "_notes" not in df.columns
    assert list(map(str, df.columns[-1:])) == ["2025"]
    assert len(df) == 2
    assert float(df.iloc[0]["2025"]) == 1.0
    assert float(df.iloc[1]["2025"]) == 2.0


def test_excel_misspelled_header_returns_partial_frame(tmp_path):
    cols = [c if c != "to key" else "to keyhhh" for c in SUPERSTRUCTURE] + ["2025"]
    path = tmp_path / "bad_header.xlsx"
    _write_xlsx(path, [cols, _sample_row(1.0)])

    df = import_from_excel(path, import_sheet=1)

    assert not df.empty
    assert is_partial_flow_sdf_headers(df.columns)
    assert missing_superstructure_columns(df.columns) == ["to key"]
    assert "to keyhhh" in df.columns


def test_excel_underscore_column_first(tmp_path):
    cols = ["_notes"] + list(SUPERSTRUCTURE) + ["2025"]
    data = ["my note"] + _sample_row(3.0)[: len(SUPERSTRUCTURE)] + [3.0]
    path = tmp_path / "sdf_first_notes.xlsx"
    _write_xlsx(path, [cols, data])

    df = import_from_excel(path, import_sheet=1)

    assert not df.empty
    assert "_notes" not in df.columns
    assert "from activity name" in df.columns
    assert float(df.iloc[0]["2025"]) == 3.0
