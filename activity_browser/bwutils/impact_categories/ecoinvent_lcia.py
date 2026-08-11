"""Ecoinvent LCIA Implementation Excel import (vendor multi-method workbook)."""
from __future__ import annotations

import functools
import warnings
from numbers import Number

import tqdm
from bw2data import Database, Method, config, methods
from bw2io.importers.base_lcia import LCIAImporter
from bw2io.strategies import (
    drop_unspecified_subcategories,
    link_iterable_by_fields,
    normalize_units,
    rationalize_method_names,
    set_biosphere_type,
)
from openpyxl import load_workbook


class EcoinventLCIAImporter(LCIAImporter):
    """Import ecoinvent-compatible LCIA Implementation Excel workbooks."""

    def __init__(self, filepath, biosphere=None):
        self.strategies = []
        self.applied_strategies = []
        self.filepath = filepath
        self.biosphere_name = biosphere
        if self.biosphere_name:
            self.set_biosphere(self.biosphere_name)

    @classmethod
    def setup_with_ei_excel(cls, file: str, biosphere_database: str | None = None):
        importer = cls(file, biosphere_database)
        importer.set_biosphere(biosphere_database or config.biosphere)
        importer.cf_data, importer.units = convert_lcia_methods_data(file)
        importer.separate_methods()
        return importer

    def set_biosphere(self, biosphere_database: str, *, relink: bool = False):
        kwargs = {"other": Database(biosphere_database), "fields": ("name", "categories")}
        if relink:
            kwargs["relink"] = True
        self.strategies = [
            normalize_units,
            set_biosphere_type,
            drop_unspecified_subcategories,
            functools.partial(link_iterable_by_fields, **kwargs),
        ]

    def add_rationalize_method_names_strategy(self):
        self.strategies.append(rationalize_method_names)

    def separate_methods(self):
        """Split flat CF rows into distinct method datasets."""
        missing = {line["method"] for line in self.cf_data if line["method"] not in self.units}
        if missing:
            warnings.warn(
                "Missing units for following: "
                + " | ".join(sorted(str(m) for m in missing))
            )

        by_method: dict[tuple, dict] = {}
        for line in self.cf_data:
            if line is None:
                continue
            assert isinstance(line["amount"], Number)
            name = line["method"]
            if name not in by_method:
                by_method[name] = {
                    "filename": self.filepath,
                    "unit": self.units.get(name, ""),
                    "name": name,
                    "description": "",
                    "exchanges": [],
                }
            by_method[name]["exchanges"].append(
                {
                    "name": line["name"],
                    "categories": line["categories"],
                    "amount": line["amount"],
                }
            )
        self.data = list(by_method.values())

    def apply_strategies(self, strategies=None, verbose=False):
        for strategy in tqdm.tqdm(
            strategies or self.strategies, desc="Applying strategies"
        ):
            self.apply_strategy(strategy, verbose=verbose)

    def prepend_methods(self, prepend: str):
        if not prepend:
            return
        for method in tqdm.tqdm(self.data, desc=f"Prepending {prepend} to ICs"):
            method["name"] = (prepend, *method["name"])

    def write_methods(self, overwrite=False, verbose=True):
        num_methods, num_cfs, num_unlinked = self.statistics(False)
        if num_unlinked:
            raise ValueError(f"Can't write unlinked methods ({num_unlinked} unlinked cfs)")
        for ds in tqdm.tqdm(self.data, desc="Writing impact categories"):
            name = ds["name"]
            if name in methods:
                if not overwrite:
                    raise ValueError(
                        f"Method {name} already exists. Use overwrite=True"
                    )
                del methods[name]
            method = Method(name)
            cfs = self._reformat_cfs(ds["exchanges"])
            method.register(
                description=ds["description"],
                filename=ds["filename"],
                unit=ds["unit"],
                num_cfs=len(cfs),
            )
            method.write(cfs)
        if verbose:
            print(
                f"Wrote {num_methods} LCIA methods with {num_cfs} characterization factors"
            )


def convert_lcia_methods_data(filename: str):
    wb = load_workbook(filename, read_only=True)

    cf_data = []
    sheet = wb["CFs"]
    for rowidx, row in tqdm.tqdm(
        enumerate(sheet.rows), total=sheet.max_row, desc="Processing CFs"
    ):
        if not rowidx:
            continue
        data = [cell.value for _, cell in zip(range(8), row)]
        if isinstance(data[-1], Number):
            cf_data.append(
                {
                    "method": tuple(data[:3]),
                    "name": data[3],
                    "categories": tuple(data[4:6]),
                    "amount": data[6],
                }
            )

    units = {}
    sheet = wb["Indicators"]
    for rowidx, row in tqdm.tqdm(
        enumerate(sheet.rows), total=sheet.max_row, desc="Processing indicators"
    ):
        if not rowidx:
            continue
        data = [cell.value for _, cell in zip(range(4), row)]
        units[tuple(data[:3])] = data[3]

    return cf_data, units
